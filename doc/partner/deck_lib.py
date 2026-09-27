"""キューテスト様向けスライド(build_*.py)で共通の部品: デザインの決まり・文字・図形・スライドの枠。

使い方:
    from deck_lib import *
    deck = Deck(total=19, footer="…")
    s = deck.new_slide("見出し", kicker="小見出し")
    text(s, 0.6, 1.5, 12, 0.5, "本文")
    deck.save(OUT)
"""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

__all__ = [
    "FONT", "DEEP", "TEAL", "MID", "TINT", "TINT2", "INK", "MUTED", "WHITE", "LINE", "SW", "SH",
    "RGBColor", "MSO_CONNECTOR", "MSO_SHAPE", "MSO_ANCHOR", "PP_ALIGN", "Inches", "Pt",
    "Deck", "jp_text", "text", "shape", "card", "badge", "arrow", "bullets", "notes",
]

# ---- デザインの決まり ---------------------------------------------------------
FONT = "IPAPGothic"  # PDF に埋め込む日本語フォント(MS Pゴシックと同じ字幅)
DEEP = RGBColor(0x0E, 0x4F, 0x5C)  # 濃いティール(表紙・見出し)
TEAL = RGBColor(0x1F, 0x8A, 0x8A)  # ティール(アイコン・強調)
MID = RGBColor(0x7C, 0xC4, 0xC0)  # 中間のティール(線・矢印)
TINT = RGBColor(0xE8, 0xF5, 0xF4)  # 薄いティール(カードの地)
TINT2 = RGBColor(0xD3, 0xEC, 0xEA)  # 少し濃い地
INK = RGBColor(0x1F, 0x2D, 0x33)  # 本文
MUTED = RGBColor(0x5B, 0x6B, 0x70)  # 補足
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC9, 0xDE, 0xDC)

SW, SH = 13.333, 7.5


def _font(run, size, color=INK, bold=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def jp_text(t):
    """
    スライドに出す文字。括弧は中の文字に合わせて全角・半角を選ぶ(開き括弧は次の文字、閉じ括弧は前の文字が日本語なら全角)。
    LibreOffice の PDF は日本語と半角の文字の間に空きを入れるため、括弧と中の文字の間に空きが入らないようにする。
    """
    out = []
    for i, c in enumerate(t):
        if c == "(" and i + 1 < len(t) and not t[i + 1].isascii():
            c = "（"
        elif c == ")" and i > 0 and not t[i - 1].isascii():
            c = "）"
        out.append(c)
    return "".join(out)


class Deck:
    """16:9 のスライド一式。ページ番号(n / total)と下の帯の文字を付ける。"""

    def __init__(self, total, footer):
        self.total = total
        self.footer = footer
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self._blank = self.prs.slide_layouts[6]
        self.page = 0
        # 後のスライドのページ番号(先のスライドから「(nページ)」で参照する。作り終えてから埋める)
        self.page_of = {}

    @staticmethod
    def page_ref(key):
        """後のスライドのページ番号の埋め込み先(fill_page_refs で埋める)。"""
        return "{page:" + key + "}"

    def fill_page_refs(self):
        for slide in self.prs.slides:
            for shp in slide.shapes:
                if not shp.has_text_frame:
                    continue
                for para in shp.text_frame.paragraphs:
                    for r in para.runs:
                        for key, n in self.page_of.items():
                            r.text = r.text.replace(self.page_ref(key), str(n))
                        assert "{page:" not in r.text, f"ページ番号を埋められません: {r.text}"

    def new_slide(self, title=None, dark=False, kicker=None):
        self.page += 1
        s = self.prs.slides.add_slide(self._blank)
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = DEEP if dark else WHITE
        if title:
            if kicker:
                text(s, 0.6, 0.38, 12, 0.4, kicker, size=16, color=TEAL, bold=True)
            text(s, 0.6, 0.62 if kicker else 0.5, 12.1, 0.8, title, size=34, color=WHITE if dark else DEEP,
                 bold=True, anchor=MSO_ANCHOR.TOP)
        if self.page > 1:
            col = RGBColor(0xB5, 0xD8, 0xD6) if dark else MUTED
            text(s, 0.6, 7.0, 9, 0.3, self.footer, size=11, color=col)
            text(s, 11.2, 7.0, 1.53, 0.3, f"{self.page} / {self.total}", size=11, color=col, align=PP_ALIGN.RIGHT)
        return s

    def save(self, out):
        assert self.page == self.total, f"TOTAL({self.total}) とスライド数({self.page})が違います"
        self.fill_page_refs()
        self.prs.save(out)
        print(f"保存しました: {out}({self.page}枚)")


def text(slide, x, y, w, h, paras, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space_after=6, line=1.15, margin=0.0):
    """paras: 文字列、または段落のリスト。段落は文字列か [(文字, {size,color,bold}), ...]。"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = anchor
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = line
        para.space_after = Pt(space_after)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, opt in runs:
            r = para.add_run()
            r.text = jp_text(t)
            _font(r, opt.get("size", size), opt.get("color", color), opt.get("bold", bold))
    return tb


def shape(slide, kind, x, y, w, h, fill=TINT, line=None, radius=0.12):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def card(slide, x, y, w, h, fill=TINT, line=None, radius=0.08):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius)


def badge(slide, x, y, d, label, fill=TEAL, color=WHITE, size=None):
    s = shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    _font(r, size or d * 30, color, True)
    return s


def arrow(slide, x1, y1, x2, y2, color=MID, width=2.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def bullets(slide, x, y, w, h, items, size=18, gap=10, mark="●", mark_color=TEAL, color=INK):
    paras = [[(mark + " ", {"color": mark_color, "size": size - 4}), (it, {})] for it in items]
    return text(slide, x, y, w, h, paras, size=size, color=color, space_after=gap, line=1.2)
